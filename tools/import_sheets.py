#!/usr/bin/env python3
"""从作者的 Google Sheets 导出（xlsx）生成空/海/导弹/防空/投送装备的 YAML 与库存 CSV。

用法：
  python3 tools/import_sheets.py --xlsx 分层级火力.xlsx
输出（可重复生成、纳入版本管理）：
  data/forcedb/equipment/{air,naval,missile_ad,transport}.yaml
  data/forcedb/inventory_cn_us.csv
约定：
  * 物理/技术参数放 specs（键名带单位后缀）。
  * 表格里已经算出来的“文章结果”放 specs 的 article.* 键（快照），引擎可在 snapshot 模式下直接使用；
    recompute 模式则用物理参数按公式重算并与快照对比。
需要 openpyxl（pip3 install openpyxl）；构建数据库本身不需要 openpyxl。
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

try:
    import openpyxl
    import yaml
except ImportError:
    sys.exit("需要 openpyxl 和 pyyaml：pip3 install openpyxl pyyaml")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "forcedb"
SRC = "sheet_cn_us"
AS_OF = "2025-08-31"

SERVICE = {"空军": "airforce", "海军": "navy", "陆军": "army", "陆战队": "marines", "海军陆战队": "marines"}

# 中文型号前缀 → 拼音/英文缩写（最长优先）。出现未覆盖的汉字会报错，补这里即可。
TOKENS = [
    ("空警", "kj"), ("歼轰", "jh"), ("歼侦", "jz"), ("歼", "j"), ("苏", "su"), ("轰", "h"), ("运油", "yy"),
    ("运干", "yg"), ("运侦", "yzh"), ("运指", "yzhi"), ("运心", "yx"), ("运潜", "yq"), ("运", "y"), ("直", "z"),
    ("红旗", "hq"), ("东风", "df"), ("长剑", "cj"), ("鹰击", "yj"), ("彩虹", "ch"), ("翼龙", "yl"),
    ("攻击", "gj"), ("无侦", "wz"), ("神雕", "sd"), ("高新", "gx"), ("米", "mi"), ("卡", "ka"), ("伊尔", "il"),
    ("轰油", "hy"), ("宽", "k"), ("式", ""), ("型", ""), ("级", ""), ("号", ""), ("伯克", "burke"),
    ("提康德罗加", "ticonderoga"), ("朱姆沃尔特", "zumwalt"), ("独立", "independence"), ("自由", "freedom"),
    ("福特", "ford"), ("尼米兹", "nimitz"), ("美国", "america"), ("黄蜂", "wasp"), ("辽宁", "liaoning"),
    ("山东", "shandong"), ("福建", "fujian"), ("现代", "sovremenny"), ("基洛", "kilo"), ("洛杉矶", "la"),
    ("海狼", "seawolf"), ("弗吉尼亚", "virginia"), ("俄亥俄", "ohio"), ("圣安东尼奥", "san_antonio"),
    ("惠德贝岛", "whidbey_island"), ("岸基", "coastal_"), ("堤丰系统", "typhon"), ("系统", ""),
    ("额外", "extra_"), ("改装", "_mod"), ("三四批次", "_b34"), ("四批次", "_b4"), ("反隐雷达", "antistealth_radar"),
    ("以前", "_early"), ("以后", "_late"), ("号舰", ""), ("中国总计", "total_cn"), ("美国总计", "total_us"),
    ("铁拳", "fist"), ("爱国者", "patriot"), ("萨德", "thaad"), ("复仇者", "avenger"), ("陆盾", "ludun"),
    ("车载红缨", "veh_hongying"), ("道尔", "tor"), ("红缨", "hongying"), ("复仇", "avenger"),
    ("弹炮合一", "gun_missile"), ("牵引", "towed"), ("自行", "sp"), ("旅", ""), ("团", ""),
    ("直升机", "heli"), ("陆战队", "usmc"), ("陆军", "army"), ("空军", "af"), ("海军", "navy"),
    ("白板", "base"), ("氮化镓", "gan"), ("涡扇", "ws"), ("野兽", "beast"), ("美军", "us"),
    ("双尾蝎", "twin_scorpion"), ("高新", "gx"), ("号", ""), ("船坞", "dock"), ("运输", "transport"), ("车载毒刺", "veh_stinger"), ("毒刺", "stinger"),
]
_TOK = sorted(TOKENS, key=lambda t: -len(t[0]))
_unresolved = set()


def slug(name: str) -> str:
    s = str(name).strip()
    for zh, en in _TOK:
        s = s.replace(zh, en)
    if re.search(r"[一-鿿]", s):
        _unresolved.add(name)
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").lower()
    return re.sub(r"_+", "_", s)


def family(eid: str) -> str:
    m = re.match(r"([a-z]+_?\d+)", eid)
    return (m.group(1).replace("_", "") if m else eid.split("_")[0])


def clean(v):
    if isinstance(v, float):
        if v != v:
            return None
        return round(v, 6) if abs(v) < 1e12 else v
    return v


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


class Sheets:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=True)

    def cell(self, sheet, ref):
        return self.wb[sheet][ref].value

    def rows(self, sheet, r0, r1):
        ws = self.wb[sheet]
        for r in range(r0, r1 + 1):
            name = ws[f"A{r}"].value
            if name is None:
                continue
            yield r, str(name).strip(), ws


class Out:
    def __init__(self):
        self.models, self.inv, self.ids = [], [], set()

    def add(self, name, category, origin, specs, *, country, service, qty=None, role="all", eid=None,
            kind="active", notes=None, confidence="stated", extra=None):
        eid = eid or slug(name)
        rec = None
        for m in self.models:
            if m["id"] == eid:
                rec = m
                break
        if rec is None:
            rec = {"id": eid, "name": name, "category": category, "family": family(eid), "origin": origin,
                   "specs": {}}
            if extra:
                rec.update(extra)
            self.models.append(rec)
        for k, v in specs.items():
            v = clean(v)
            if v is not None:
                rec["specs"][k] = v
        if notes:
            rec["notes"] = notes
        if qty is not None:
            self.inv.append(dict(country=country, service=service, equipment_id=eid, kind=kind, role=role,
                                 quantity=clean(qty), as_of=AS_OF, source_id=SRC, confidence=confidence, notes=""))
        return eid


def sp(ws, r, mapping):
    return {k: ws[f"{c}{r}"].value for k, c in mapping.items()}


# ------------------------------------------------------------------ 航空
def import_air(S: Sheets, o: Out):
    # 战斗机 + 战斗机对地（按行对齐）
    fighters = {
        "thrust_weight": "D", "max_speed_mach": "E", "cruise_speed_mach": "F", "combat_radius_km": "G",
        "aam_count": "H", "aam_range_km": "I", "nose_diameter_mm": "J", "radar_type": "K",
        "radar_range_mult": "L", "detect_rcs3_km": "M", "rcs_m2": "N", "detected_by_j16_km": "O",
        "article.air_power": "P", "article.air_power_eq_fa18": "Q", "article.air_total": "R",
    }
    ground = {
        "ground_radius_km": "E", "ag_payload_t": "F", "ground_detected_km": "H",
        "article.ground_power": "I", "article.ground_power_eq_fa18": "J", "article.ground_total": "K",
    }
    for (c0, c1, country, origin_hint) in ((2, 28, "CHN", "CHN"), (33, 48, "USA", "USA")):
        wsA, wsG = S.wb["战斗机"], S.wb["战斗机对地"]
        for r in range(c0, c1 + 1):
            name = wsA[f"A{r}"].value
            if name is None:
                continue
            if wsG[f"A{r}"].value != name:
                raise ValueError(f"战斗机/对地 行未对齐：{r} {name} vs {wsG[f'A{r}'].value}")
            origin = "RUS" if name.startswith("苏") else origin_hint
            svc = SERVICE[wsA[f"B{r}"].value]
            eid = slug(name)
            if name == "F-35C":   # 表里 F-35C 分别出现在海军与陆战队两行，参数不同，分开建模
                eid = "f35c_usmc" if svc == "marines" else "f35c"
            if name in ("F/A-18C/D氮化镓", "F/A-18C/D") and svc == "marines":
                eid = slug(name)
            specs = sp(wsA, r, fighters)
            specs.update(sp(wsG, r, ground))
            specs["ground_cruise_mach"] = wsG[f"D{r}"].value
            specs["ground_rcs_m2"] = wsG[f"G{r}"].value
            o.add(name, "fighter", origin, specs, country=country, service=svc, qty=wsA[f"C{r}"].value,
                  role="air_superiority", eid=eid)
            o.inv.append(dict(country=country, service=svc, equipment_id=eid, kind="active", role="ground_attack",
                              quantity=clean(wsG[f"C{r}"].value), as_of=AS_OF, source_id=SRC, confidence="stated",
                              notes=""))
    # 轰炸攻击机
    bm = {"cruise_speed_mach": "C", "combat_radius_km": "D", "ag_payload_t": "E", "rcs_m2": "F",
          "detected_by_j16_km": "G", "article.ground_power": "H", "article.ground_power_eq_fa18": "I",
          "article.ground_total": "J"}
    for (rr, country) in ((range(2, 4), "CHN"), (range(6, 11), "USA")):
        for r in rr:
            ws = S.wb["轰炸攻击机"]
            name = ws[f"A{r}"].value
            origin = country
            o.add(name, "bomber", origin, sp(ws, r, bm), country=country, service="airforce",
                  qty=ws[f"B{r}"].value, role="ground_attack")
    # 预警机
    am = {"payload_t": "C", "endurance_h": "D", "antenna_form": "E", "antenna_diameter_m": "F",
          "antenna_thickness_m": "G", "antenna_area_m2": "H", "radar_type": "I", "radar_range_mult": "J",
          "detect_rcs3_km": "K", "article.composite": "L", "article.power_eq_fa18": "M", "article.power_total": "N"}
    for rr, country in ((range(2, 8), "CHN"), (range(12, 15), "USA")):
        for r in rr:
            ws = S.wb["预警机"]
            name = ws[f"A{r}"].value
            svc = "navy" if name in ("空警-600", "E-2C", "E-2D") else "airforce"
            o.add(name, "awacs", country, sp(ws, r, am), country=country, service=svc, qty=ws[f"B{r}"].value)
    # 加油机
    tm = {"fuel_capacity_t": "C", "transferable_fuel_t": "D", "article.composite": "E",
          "article.power_eq_fa18": "F", "article.power_total": "G"}
    for rr, country in ((range(2, 5), "CHN"), (range(7, 11), "USA")):
        for r in rr:
            ws = S.wb["加油机"]
            name = ws[f"A{r}"].value
            origin = "RUS" if name.startswith("伊尔") else country
            svc = "navy" if name == "KC-130T" else "airforce"
            o.add(name, "tanker_aircraft", origin, sp(ws, r, tm), country=country, service=svc,
                  qty=ws[f"B{r}"].value)
    # 特种机
    xm = {"function": "C", "payload_t": "E", "range_km": "F", "unit_performance": "G", "unit_cost_musd": "H",
          "article.conv_factor": "I", "article.power_eq_fa18": "J", "article.power_total": "K"}
    for rr, country in ((range(2, 17), "CHN"), (range(26, 40), "USA")):
        for r in rr:
            ws = S.wb["特种机"]
            name = ws[f"A{r}"].value
            cat = "asw_aircraft" if ws[f"C{r}"].value == "反潜" else "ew_aircraft"
            svc = SERVICE[ws[f"B{r}"].value]
            o.add(name, cat, country, sp(ws, r, xm), country=country, service=svc, qty=ws[f"D{r}"].value)
    # 无人机
    um = {"payload_t": "D", "mtow_t": "E", "endurance_h": "F", "max_speed_kph": "G", "range_km": "H",
          "stealth_factor": "I", "article.composite": "J", "article.power_eq_fa18": "K", "article.power_total": "L"}
    for rr, country in ((range(2, 14), "CHN"), (range(19, 27), "USA")):
        for r in rr:
            ws = S.wb["无人机"]
            name = ws[f"A{r}"].value
            svc = SERVICE[ws[f"B{r}"].value]
            eid = slug(name)
            o.add(name, "large_uav", country, sp(ws, r, um), country=country, service=svc,
                  qty=ws[f"C{r}"].value, eid=eid)


# ------------------------------------------------------------------ 海军
SHIP_CAT = {
    "伯克": "destroyer", "提康德罗加": "cruiser", "朱姆沃尔特": "destroyer", "独立级": "corvette", "自由级": "corvette",
    "052": "destroyer", "051": "destroyer", "现代级": "destroyer", "053": "frigate", "054": "frigate",
    "055": "cruiser", "056": "corvette", "022": "corvette",
}
WEAPON_COLS = {"cell_htk1_deep": "C", "cell_htk1": "D", "cell_mk41": "E", "cell_sea_ram_riff": "F", "cell_hq9": "G",
               "cell_hq16": "H", "cell_shtkili": "I", "antiship_missiles": "J", "point_defense": "K",
               "asw_rocket": "L", "torpedoes": "M"}


def ship_category(name):
    for k, v in SHIP_CAT.items():
        if k in name:
            return v
    return "destroyer"


def import_naval(S: Sheets, o: Out):
    ws = S.wb["驱护舰"]
    coef = {k: ws[f"{c}2"].value for k, c in WEAPON_COLS.items()}
    o.meta = {"ship_weapon_coefficients": coef}
    ship_m = {"radar_height_m": "O", "radar_width_m": "P", "radar_faces": "Q", "radar_area_m2": "R",
              "radar_type": "S", "radar_range_mult": "T", "detect_rcs3_km": "U", "article.radar_vls": "V",
              "article.firepower_vls": "N", "article.power_per_ship": "W"}
    for rr, country in ((list(range(3, 14)), "USA"), (list(range(16, 34)), "CHN")):
        for r in rr:
            name = ws[f"A{r}"].value
            if name is None:
                continue
            coastal = name.startswith("岸基")
            cat = "coastal_missile_battery" if coastal else ship_category(name)
            specs = sp(ws, r, ship_m)
            for k, c in WEAPON_COLS.items():
                specs["weapons." + k] = ws[f"{c}{r}"].value
            eid = slug(name)
            if name == "伯克额外SeaRAM":
                cat = "naval_sam"
            origin = "RUS" if name.startswith("现代级") else country
            o.add(name, cat, origin, specs, country=country, service="navy", qty=ws[f"B{r}"].value, eid=eid)
    # 潜艇
    ws = S.wb["潜艇"]
    sm = {"displacement_t": "C", "noise_db": "D", "detect_km": "E", "propulsion": "F", "endurance_5kn_days": "G",
          "endurance_15kn_days": "H", "tubes_vls": "I", "weapons_loaded": "J", "article.power_per_boat": "K",
          "article.power_total": "L"}
    sub_cat = {"洛杉矶级": "ssn", "海狼级": "ssn", "弗吉尼亚级": "ssn", "俄亥俄级": "ssg_n",
               "093": "ssn", "093A": "ssn", "093B": "ssn"}
    for rr, country in ((range(2, 6), "USA"), (range(9, 18), "CHN")):
        for r in rr:
            name = ws[f"A{r}"].value
            cat = sub_cat.get(name, "ssk")
            origin = "RUS" if name in ("基洛级",) else country
            eid = "sub_" + slug(name)
            o.add(name, cat, origin, sp(ws, r, sm), country=country, service="navy", qty=ws[f"B{r}"].value, eid=eid)
    # 航母 / 两栖攻击舰
    ws = S.wb["航母"]
    cm = {"sorties_fixed_wing_day": "C", "sorties_heli_day": "D", "article.sorties_eq": "E", "article.power_vls": "G"}
    carrier_cat = {"福特级": "aircraft_carrier", "尼米兹级": "aircraft_carrier", "美国级": "amphibious_carrier",
                   "黄蜂级": "amphibious_carrier", "辽宁/山东": "aircraft_carrier", "福建": "aircraft_carrier",
                   "075": "amphibious_carrier"}
    for rr, country in ((range(2, 6), "USA"), (range(9, 12), "CHN")):
        for r in rr:
            name = ws[f"A{r}"].value
            eid = "cv_" + slug(name) if re.match(r"^\d", slug(name)) else slug(name)
            o.add(name, carrier_cat[name], country, sp(ws, r, cm), country=country, service="navy",
                  qty=ws[f"B{r}"].value, eid=eid)


# ------------------------------------------------------------------ 投送
def import_transport(S: Sheets, o: Out):
    ws = S.wb["运输机"]
    m = {"payload_t": "C", "full_load_range_km": "E", "ferry_range_km": "F", "cruise_speed_kph": "G",
         "turnaround_prep_h": "I", "turnaround_load_h": "J", "takeoff_landing_h": "K",
         "article.rtt_total_h": "N", "article.throughput_24h_t": "O"}
    for rr, country in ((range(2, 8), "USA"), (range(10, 15), "CHN")):
        for r in rr:
            name = ws[f"A{r}"].value
            origin = "RUS" if name.startswith("伊尔") else country
            qty = ws[f"B{r}"].value
            if isinstance(qty, str):
                qty = None
            eid = "airlift_" + slug(name)
            qv = qty if qty is not None else eval_simple(ws, f"B{r}", S)
            o.add(name, "transport_aircraft", origin, sp(ws, r, m), country=country,
                  service="airforce", qty=qv, eid=eid)
    ws = S.wb["运输舰"]
    m = {"displacement_t": "B", "payload_t": "C", "max_speed_kn": "F", "econ_speed_kn": "G", "range_nm": "H",
         "load_h": "I", "port_h": "J", "unload_h": "K", "article.rtt_total_h": "N", "article.throughput_24h_t": "O"}
    for rr, country in ((range(2, 6), "USA"), (range(8, 19), "CHN")):
        for r in rr:
            name = ws[f"A{r}"].value
            cat = "amphibious_dock" if country == "USA" or name in ("075", "071") else "logistics_ship"
            eid = "sealift_" + slug(name)
            o.add(name, cat, country, sp(ws, r, m), country=country, service="navy", qty=ws[f"D{r}"].value, eid=eid)


def eval_simple(ws, ref, S):
    # data_only 读到的是值；C-130 的数量公式单元格 data_only 即为数值
    v = ws[ref].value
    return v if isinstance(v, (int, float)) else None


# ------------------------------------------------------------------ 导弹 / 炮兵 / 防空
def import_missile_ad(S: Sheets, o: Out):
    ws = S.wb["火箭军"]
    m = {"warhead_kg": "B", "missiles_per_launcher": "C", "range_km": "D", "cep_m": "E", "article.base_param": "F",
         "article.mobility": "G", "article.reload": "H", "article.power": "I"}
    cat_by_range = lambda km: ("srbm" if km <= 1000 else "mrbm" if km <= 3000 else "irbm" if km <= 5500 else "icbm")
    for r in range(2, 15):
        name = ws[f"A{r}"].value
        km = ws[f"D{r}"].value
        cat = "cruise_missile_land" if name.startswith("长剑") else cat_by_range(km)
        eid = slug(name)
        specs = sp(ws, r, m)
        specs["brigades"] = ws[f"J{r}"].value
        specs["launchers_per_brigade"] = ws[f"K{r}"].value
        launchers = (ws[f"J{r}"].value or 0) * (ws[f"K{r}"].value or 0)
        o.add(name, cat, "CHN", specs, country="CHN", service="rocket", qty=launchers, eid=eid,
              extra={"notes": "库存口径为发射车数 = 旅数 × 每旅发射车数（表格“火箭军”）。"})
    # 火炮折算（炮兵）
    ws = S.wb["火炮"]
    am = {"barrels": "B", "guided_warhead_kg": "C", "guided_range_km": "D", "cep_m": "E", "article.base_param": "F",
          "deploy_speed": "G", "traverse": "H", "aiming": "I", "mobility": "J", "reload": "K", "article.power_per_gun": "L"}
    art_cat = {"M777": "howitzer_heavy_towed", "M109A6": "howitzer_heavy_sp", "PLZ-05": "howitzer_heavy_sp"}
    for r in range(2, 16):
        name = ws[f"A{r}"].value
        origin = "CHN" if name.startswith(("PLZ", "PCL", "PHL")) else "USA"
        cat = art_cat.get(name, "mrl_long_range" if "PHL" in name or "M270" in name or "HIMARS" in name else
                          "howitzer_heavy_truck" if name.startswith("PCL") else "mrl_long_range")
        eid = "art_" + slug(name)
        o.add(name, cat, origin, sp(ws, r, am), country=origin, service="army", eid=eid)
    # 炮兵旅数量（等效营数 × 每营火炮 → 库存口径 = 营数×每营门数）
    ws = S.wb["炮兵旅"]
    for r, country in ((2, "USA"), (3, "USA"), (4, "USA"), (5, "USA"), (9, "CHN"), (10, "CHN"), (11, "CHN"), (12, "CHN")):
        name = ws[f"A{r}"].value
        eid = "art_" + slug(name)
        battalions, per_bn = ws[f"B{r}"].value, ws[f"C{r}"].value
        if eid not in {m["id"] for m in o.models}:   # HIMARS/M270/PHL-191 在“炮兵旅”表按系统汇总，各弹种在“火炮”表
            o.add(name, "mrl_long_range", country, {"article.power_per_gun": ws[f"D{r}"].value},
                  country=country, service="army", eid=eid,
                  extra={"notes": "系统级条目：各弹种配置见 art_<系统>_<弹种>；战力/门取自“炮兵旅”表。"})
        o.inv.append(dict(country=country, service="army", equipment_id=eid, kind="in_toe", role="article_equivalent",
                          quantity=battalions * per_bn, as_of=AS_OF, source_id=SRC, confidence="estimated",
                          notes="等效营数×每营门数（国民警卫队对半折算，见 orbat）"))
    # 防空
    ws = S.wb["防空"]
    dm = {"range_km": "B", "ceiling_km": "C", "missile_speed_mach": "D", "gun_rof": "E", "guidance": "F",
          "guidance_factor": "G", "deploy_min": "H", "deploy_factor": "I", "launchers": "J",
          "article.battery_power": "K", "article.brigade_equiv": "L", "vehicles_per_unit": "N"}
    for rr, country in ((range(2, 24), "CHN"), (range(27, 36), "USA")):
        for r in rr:
            name = ws[f"A{r}"].value
            if name is None:
                continue
            units = ws[f"M{r}"].value
            origin = "RUS" if name in ("S400", "S300PMU2", "S300PMU1", "S300PMU", "道尔M1") else country
            cat = ("sam_long" if (ws[f"B{r}"].value or 0) >= 100 else "sam_medium" if (ws[f"B{r}"].value or 0) >= 20
                   else "sam_short" if ws[f"E{r}"].value is None else "spaag")
            if name.startswith(("PGL", "PG-", "陆盾", "C-RAM", "M-Shorad")) and ws[f"E{r}"].value:
                cat = "spaag"
            svc = "airforce" if (country == "CHN" and r <= 13 or name == "陆盾2000") else "army"
            if name == "萨德" or name.startswith("爱国者") or name in ("复仇者",):
                svc = "army"
            eid = "ad_" + slug(name)
            o.add(name, cat, origin, sp(ws, r, dm), country=country, service=svc, qty=units, eid=eid,
                  extra={"notes": "库存口径 = 营/连数（表格“防空”M 列）；每单位发射车数见 vehicles_per_unit。"})


# ------------------------------------------------------------------ 陆军：旅级单位 / 直升机 / 特战
BRIG_ID = {
    "ABCT": "us_abct", "ABCT国民警卫队": "us_abct_ng", "第二骑兵团": "us_2cr_sbct", "第二步兵师第二旅": "us_2id_2bct_sbct",
    "其他SBCT": "us_sbct_other", "SBCT国民警卫队": "us_sbct_ng", "IBCT": "us_ibct", "空突IBCT": "us_air_assault_ibct",
    "空降IBCT": "us_airborne_ibct", "国民警卫队IBCT": "us_ibct_ng", "陆战队MEU": "us_marine_meu",
    "陆战队MEU预备役": "us_marine_meu_reserve",
    "99A+04A旅": "cn_99a_04a_bde", "99+04A旅": "cn_99_04a_bde", "99+86A旅": "cn_99_86a_bde", "15+04A旅": "cn_15_04a_bde",
    "15+86A旅": "cn_15_86a_bde", "96A+04A旅": "cn_96a_04a_bde", "96A+04旅": "cn_96a_04_bde",
    "96A+86A旅": "cn_96a_86a_bde", "96+86A旅": "cn_96_86a_bde", "96+86旅": "cn_96_86_bde", "96+89旅": "cn_96_89_bde",
    "88B+89A旅": "cn_88b_89a_bde", "两栖旅": "cn_amphibious_bde", "191车族旅": "cn_191_family_bde",
    "11+08旅": "cn_11_08_bde", "11+92旅": "cn_11_92_bde", "猛士三代旅": "cn_hmv3_bde", "猛士二代旅": "cn_hmv2_bde",
    "山地旅": "cn_mountain_bde", "摩步旅": "cn_motor_bde", "96A+86A团": "cn_96a_86a_regt", "11+08团": "cn_11_08_regt",
    "三代猛士团": "cn_hmv3_regt", "摩步团": "cn_motor_regt", "空降履带旅": "cn_abn_tracked_bde",
    "空降轮式旅": "cn_abn_wheeled_bde", "空降山猫旅": "cn_abn_lynx_bde", "空降摩步旅": "cn_abn_motor_bde",
    "空突山猫旅": "cn_air_assault_lynx_bde", "空突摩步旅": "cn_air_assault_motor_bde", "海防旅": "cn_coastal_bde",
    "边防团": "cn_border_regt", "边防旅": "cn_border_bde", "驻港合成旅": "cn_hk_garrison_bde",
}
AIRFORCE_CN_ROWS = {40, 41, 42, 43, 45}     # 表格“陆军!B4”只合计这五行（空突山猫旅第 44 行并入陆军，原样保留）


def import_brigades(S: Sheets):
    ws = S.wb["乘数"]
    rows = []
    spec = [(range(2, 14), "USA"), (range(16, 51), "CHN")]
    for rr, country in spec:
        for r in rr:
            label = ws[f"A{r}"].value
            if label is None:
                continue
            b = ws[f"B{r}"].value
            if label == "陆战队合成旅":   # 表格里该行只有汇总战力（由两栖旅/空突/11+08 旅平均估算）
                rows.append(dict(id="cn_marine_brigade", country=country, label=label, service="marines", count=b,
                                 base_fire_scale="", equip_mult="", mech_mult="", info_mult="", logistics_mult="",
                                 readiness_mult="", apply_logistics="", snapshot_power=ws[f"I{r}"].value,
                                 snapshot_total=ws[f"J{r}"].value))
                continue
            svc = "marines" if country == "USA" and r in (12, 13) else \
                  "airforce" if country == "CHN" and r in AIRFORCE_CN_ROWS else "army"
            is_regiment = country == "CHN" and r in (36, 37, 38, 39)   # 表格中团级行的 I 列不乘后勤乘数
            rows.append(dict(
                id=BRIG_ID[label], country=country, label=label, service=svc,
                count=b, base_fire_scale=ws[f"C{r}"].value, equip_mult=ws[f"D{r}"].value,
                mech_mult=ws[f"E{r}"].value, info_mult=ws[f"F{r}"].value, logistics_mult=ws[f"G{r}"].value,
                readiness_mult=ws[f"H{r}"].value, apply_logistics=("0" if is_regiment or country == "USA" else "1"),
                snapshot_power=ws[f"I{r}"].value, snapshot_total=ws[f"J{r}"].value))
    (OUT / "toe").mkdir(parents=True, exist_ok=True)
    cols = list(rows[0].keys())
    with open(OUT / "toe" / "brigade_classes.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if v is None else clean(v)) for k, v in r.items()})
    return len(rows)


HELI_ID = {"直9": "z9", "卡-28": "ka28", "直20": "z20", "米-171": "mi171", "直8/直18": "z8", "直8宽": "z8l",
           "直10": "z10", "直19": "z19", "UH-1": "uh1y", "UH-60": "uh60", "V-22": "mv22", "CH-47D": "ch47f",
           "CH-53E": "ch53e", "AH-64": "ah64e", "AH-1Z": "ah1z", "MH-6": "mh6m"}
HELI_COLS = [("D", "army", "active"), ("E", "army", "reserve"), ("F", "army", "active"), ("G", "airforce", "active"),
             ("H", "airforce", "active"), ("I", "marines", "active"), ("J", "marines", "reserve"),
             ("K", "navy", "active")]
HELI_ROLE = {"D": "field", "E": "national_guard", "F": "sof", "G": "airborne_corps", "H": "hk_garrison",
             "I": "marines", "J": "marines_reserve", "K": "shipborne"}


def import_helicopters(S: Sheets, o: Out):
    ws = S.wb["直升机"]
    ann = []
    for rr, country in ((range(2, 10), "CHN"), (range(12, 20), "USA")):
        for r in rr:
            name = ws[f"A{r}"].value
            eid = HELI_ID[name]
            ann.append({"equipment_id": eid, "specs": {"article.power_factor": ws[f"B{r}"].value,
                                                         "article.brigade_equiv_per_airframe": ws[f"C{r}"].value}})
            for col, svc, kind in HELI_COLS:
                q = ws[f"{col}{r}"].value
                if isinstance(q, (int, float)) and q:
                    o.inv.append(dict(country=country, service=svc, equipment_id=eid, kind=kind,
                                      role=HELI_ROLE[col], quantity=clean(q), as_of=AS_OF, source_id=SRC,
                                      confidence="stated", notes=""))
    return ann


def special_forces_snapshot(S: Sheets):
    """特战：表格里 Tier 加权后的结果（按军种分），作为快照常量。"""
    ws = S.wb["特战"]
    g = lambda ref: clean(ws[ref].value)
    return {"CHN": {"total": g("G40"), "army": g("H40"), "navy": g("J40"), "airforce": g("K40")},
            "USA": {"total": g("G21"), "army": g("H21"), "marines": g("I21"), "navy": g("J21"), "airforce": g("K21")}}


# ------------------------------------------------------------------ 输出
def write(o: Out, name: str, filt):
    path = OUT / "equipment" / f"{name}.yaml"
    path.parent.mkdir(parents=True, exist_ok=True)
    models = [m for m in o.models if filt(m)]
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# 由 tools/import_sheets.py 从作者的 Google Sheets 导出生成（来源 {SRC}）。\n"
                "# 物理/技术参数在 specs；article.* 为表格中已算好的文章结果（快照）。\n")
        yaml.safe_dump({"models": models}, f, allow_unicode=True, sort_keys=False, width=120)
    return len(models)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, help="分层级火力 工作簿（xlsx 导出）")
    args = ap.parse_args()
    S = Sheets(args.xlsx)
    o = Out()
    import_air(S, o)
    import_naval(S, o)
    import_transport(S, o)
    import_missile_ad(S, o)
    n_brig = import_brigades(S)
    heli_ann = import_helicopters(S, o)
    with open(OUT / "spec_annotations.yaml", "w", encoding="utf-8") as f:
        f.write("# 为既有装备（仓库 data/equipment）追加量化引擎需要的规格/文章快照。由 tools/import_sheets.py 生成。\n")
        yaml.safe_dump({"annotations": heli_ann}, f, allow_unicode=True, sort_keys=False)
    with open(OUT / "special_forces_snapshot.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(special_forces_snapshot(S), f, allow_unicode=True, sort_keys=False)
    if _unresolved:
        sys.exit("未能转写的型号（请在 TOKENS 补充）：" + "、".join(sorted(_unresolved)))
    ids = [m["id"] for m in o.models]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        sys.exit(f"id 重复：{dup}")
    domain_of = {"fighter": "air", "bomber": "air", "awacs": "air", "tanker_aircraft": "air", "ew_aircraft": "air",
                 "asw_aircraft": "air", "large_uav": "air"}
    n_air = write(o, "air", lambda m: m["category"] in domain_of)
    n_sea = write(o, "naval", lambda m: m["category"] in ("destroyer", "cruiser", "corvette", "frigate", "ssn", "ssg_n",
                                                          "ssk", "aircraft_carrier", "amphibious_carrier",
                                                          "coastal_missile_battery", "naval_sam"))
    n_tr = write(o, "transport", lambda m: m["category"] in ("transport_aircraft", "amphibious_dock", "logistics_ship"))
    n_ms = write(o, "missile_ad", lambda m: m["category"] in ("srbm", "mrbm", "irbm", "icbm", "cruise_missile_land",
                                                            "howitzer_heavy_towed", "howitzer_heavy_sp",
                                                            "howitzer_heavy_truck", "mrl_long_range", "sam_long",
                                                            "sam_medium", "sam_short", "spaag"))
    inv_path = OUT / "inventory_cn_us.csv"
    cols = ["country", "service", "equipment_id", "kind", "role", "quantity", "as_of", "source_id", "confidence", "notes"]
    with open(inv_path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in o.inv:
            if r["quantity"] is not None:
                w.writerow(r)
    meta = OUT / "naval_coefficients.yaml"
    with open(meta, "w", encoding="utf-8") as f:
        yaml.safe_dump(o.meta, f, allow_unicode=True, sort_keys=False)
    print(f"旅级单位 {n_brig}；air {n_air}, naval {n_sea}, transport {n_tr}, missile/ad/artillery {n_ms}; 库存行 {len(o.inv)}")


if __name__ == "__main__":
    main()
